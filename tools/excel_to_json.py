"""Prevedie Excel 'kalorie - ja' do JSON zálohy, ktorú vieš importovať v appke (Nastavenia -> Import).

Použitie:
    python tools/excel_to_json.py "cesta/k/kalorie.xlsx" [vystup.json]
"""
import json
import re
import sys
from datetime import datetime, timedelta

import openpyxl


def num(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def clean(d):
    return {k: (int(v) if isinstance(v, float) and v.is_integer() else v) for k, v in d.items() if v not in (None, 0)}


def main(src, out):
    wb = openpyxl.load_workbook(src, data_only=True)
    wf = openpyxl.load_workbook(src)  # kvôli koeficientu krokov zo vzorca
    ciel, suhrn = wb["ciel"], wb["suhrn"]
    zapis, makra = wb["zapis_minusovych_kalorii"], wb["zapis makra"]

    days = {}

    # zapis_minusovych_kalorii: A dátum, B posilka, C športovanie, D kroky, G váha
    start = zapis["A2"].value.date()
    step_coef = []  # [{from, value}] - koeficient sa v Exceli menil v čase
    for r in range(2, zapis.max_row + 1):
        d = zapis.cell(r, 1).value
        if not isinstance(d, datetime):
            continue
        m = re.search(r"=\s*([\d.]+)\s*\*", str(wf["zapis_minusovych_kalorii"].cell(r, 5).value))
        if m and (not step_coef or step_coef[-1]["value"] != float(m.group(1))):
            step_coef.append({"from": d.date().isoformat(), "value": float(m.group(1))})
        days.setdefault(d.date().isoformat(), {}).update(
            gym=num(zapis.cell(r, 2).value),
            sport=num(zapis.cell(r, 3).value),
            steps=num(zapis.cell(r, 4).value),
            weight=num(zapis.cell(r, 7).value),
        )

    # suhrn: riadok = týždeň (od riadku 3), stĺpce B..H = Po..Ne zjedené kcal
    for r in range(3, suhrn.max_row + 1):
        for k in range(7):
            v = num(suhrn.cell(r, 2 + k).value)
            if v:
                d = start + timedelta(days=(r - 3) * 7 + k)
                days.setdefault(d.isoformat(), {})["kcal"] = v

    # zapis makra: A dátum, B sacharidy, C bielkoviny, D tuky, E vláknina; N..Q ciele
    targets = None
    for r in range(3, makra.max_row + 1):
        d = makra.cell(r, 1).value
        if not isinstance(d, datetime):
            continue
        if targets is None and num(makra.cell(r, 14).value):
            targets = {
                "carbs": num(makra.cell(r, 14).value),
                "protein": num(makra.cell(r, 15).value),
                "fat": num(makra.cell(r, 16).value),
                "fiber": num(makra.cell(r, 17).value),
            }
        days.setdefault(d.date().isoformat(), {}).update(
            carbs=num(makra.cell(r, 2).value),
            protein=num(makra.cell(r, 3).value),
            fat=num(makra.cell(r, 4).value),
            fiber=num(makra.cell(r, 5).value),
        )

    # Dni bez akýchkoľvek dát vyhodíme; váhu nechávame len pri dňoch, kde je niečo zapísané
    out_days = {}
    for date, d in sorted(days.items()):
        has_data = any(d.get(k) for k in ("kcal", "gym", "sport", "steps", "carbs", "protein", "fat", "fiber"))
        if has_data:
            out_days[date] = clean(d)

    # merania (ciel T:W od riadku 2)
    measurements = []
    for r in range(2, 50):
        dt = ciel.cell(r, 20).value
        if dt is None:
            continue
        if isinstance(dt, datetime):
            iso = dt.date().isoformat()
        else:
            iso = datetime.strptime(str(dt).strip(), "%d.%m.%Y").date().isoformat()
        measurements.append(clean({
            "date": iso,
            "bazal": num(ciel.cell(r, 21).value),
            "fatKg": num(ciel.cell(r, 22).value),
            "muscleKg": num(ciel.cell(r, 23).value),
            "note": ciel.cell(r, 24).value,
        }))

    data = {
        "app": "makra",
        "version": 1,
        "exported": datetime.now().isoformat(timespec="seconds"),
        "settings": {
            "startDate": start.isoformat(),
            "bazal": [{"from": start.isoformat(), "value": num(ciel["C16"].value) or 1980}],
            "startWeight": num(ciel["C17"].value) or 100,
            "goalFatKg": num(ciel["C18"].value) or 10,
            "stepCoef": step_coef or [{"from": start.isoformat(), "value": 0.000595}],
            "targets": targets or {"carbs": 110, "protein": 140, "fat": 70, "fiber": 25},
        },
        "days": out_days,
        "measurements": measurements,
    }
    with open(out, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    print(f"OK: {len(out_days)} dní, {len(measurements)} meraní -> {out}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else "makra-import.json")
